"""
Self-check for the Exports module.

The dates are where a wrong answer costs money: a deadline a month late is a
bill the bank reports as overdue, so each regime boundary is pinned. Payments
are the other: nothing may be set against an invoice until the user says so.

Imported into selftest.py, so `python -m app.backend.selftest` runs it too.
"""

from datetime import date

D = date.fromisoformat


def _sb(inv_no, sb_no, sb_date, value, currency="USD", buyer="MAPLE HOME GOODS INC"):
    return {"id": f"sb{sb_no}", "name": f"{sb_no}.pdf", "type": "sb", "port": "INJAI6",
            "sb_no": sb_no, "sb_date": sb_date, "gstin": "08AMHCK7032W1ZX", "mode": "LAND",
            "containers": [], "buyer": buyer, "exporter": "KESARI TEXTILE MILLS",
            "invoices": [{"inv_no": inv_no, "inv_date": sb_date, "currency": currency,
                          "value_fc": value, "fob_fc": value, "rate": 83.0, "terms": "FOB",
                          "igst_amount": 0.0}]}


def _advice(amount, credit_date, remitter="MAPLE HOME GOODS", currency="USD"):
    return {"id": "f1", "name": "advice.pdf", "type": "fira", "ref": "FIRA0001",
            "credit_date": credit_date, "currency": currency, "amount_fc": amount,
            "remitter": remitter, "purpose_code": "P0102", "advance": False, "invoice_ref": None}


def test_export_realisation_periods():
    from app.backend.exim import add_months, due_date, load_rules
    rules = load_rules()

    def months(day, cur="USD"):
        return due_date(D(day), cur, rules)[1]

    assert (months("2025-11-13"), months("2025-11-14")) == (9, 15)
    assert (months("2026-06-04"), months("2026-06-05")) == (15, 9)
    assert (months("2026-09-30"), months("2026-10-01")) == (9, 15)
    assert (months("2026-10-01", "INR"), months("2026-09-30", "INR")) == (18, 9)
    assert add_months(D("2025-11-30"), 15) == D("2027-02-28")
    assert add_months(D("2026-11-30"), 15) == D("2028-02-29")


def test_export_payment_waits_for_a_click():
    """A likely payment is only a question until confirmed; confirming it pays
    the invoice (bank charges allowed for) and puts it in the eBRC bulk sheet."""
    from app.backend import exim
    docs = [_sb("KE/26-27/101", "7100001", "2026-07-01", 10000.0), _advice(9980.0, "2026-08-10")]
    got = exim.build(docs, today=D("2027-02-01"))            # 9 months: due 1 Apr 2027
    by = {r["status"]: r for r in got["all_records"]}
    row = got["invoices"][0]
    assert not row["receipts"], "a payment was matched without a click"
    assert row["due"]["date"] == "2027-04-01" and row["due"]["state"] == "due_soon"
    assert got["figures"]["due_90_count"] == 1
    ask = by["pay_confirm"]
    assert ask["counterpart"] == {"kind": "payment", "remittance": ask["counterpart"]["remittance"],
                                  "invoice": row["key"], "amount": 9980.0}

    refused = exim.build(docs, {ask["fp"]: {"action": "ignore"}}, today=D("2027-02-01"))
    statuses = {r["status"] for r in refused["all_records"]}
    assert "pay_unplaced" in statuses and "pay_confirm" not in statuses, "a refused pairing came back"

    decisions = {ask["fp"]: {"action": "match", "context": ask["counterpart"]}}
    got = exim.build(docs, decisions, today=D("2027-02-01"))
    row = got["invoices"][0]
    assert [m["how"] for m in row["receipts"]] == ["you confirmed"]
    assert row["due"]["state"] == "paid" and row["due"]["short_by_charges"] == 20.0
    statuses = {r["status"] for r in got["all_records"]}
    assert statuses == {"ebrc_due", "no_bl"}, statuses
    assert "add the DGFT IRM list" in got["bulk_rows"][0]["problems"][0]


def test_export_payment_split_over_invoices():
    """Three invoices of EUR 10,000 and one advice of EUR 25,000: the settle
    window suggests the two oldest in full and 5,000 on the third, refuses a
    split that asks too much, and once saved the split pays exactly that - over
    whatever the advice names - with nothing else offered for the money."""
    from app.backend import exim
    sbs = [_sb(f"KE/26-27/10{i}", f"720000{i}", f"2026-07-0{i}", 10000.0, "EUR", "CASA OLIVA SRL") for i in (3, 1, 2)]
    advice = _advice(25000.0, "2026-09-01", remitter="CASA OLIVA S.R.L.", currency="EUR")
    today = D("2026-09-20")
    view = exim.build(sbs + [advice], today=today)
    rem_id = view["remittances"][0]["id"]
    o = exim.settle_options(view, rem_id)
    assert [(c["inv_no"], c["suggest"]) for c in o["candidates"]] == [
        ("KE/26-27/101", 10000.0), ("KE/26-27/102", 10000.0), ("KE/26-27/103", 5000.0)], o["candidates"]
    first = o["candidates"][0]["key"]
    for bad, why in (([{"invoice": first, "amount": 12000}], "left to come"),
                     ([{"invoice": c["key"], "amount": 10000} for c in o["candidates"]], "to place"),
                     ([{"invoice": "NOT-ON-FILE", "amount": 10}], "cannot take"),
                     ([{"invoice": first, "amount": -1}], "zero or more"),
                     ([{"invoice": first, "amount": "10000"}], "zero or more")):
        try:
            exim.check_splits(o, bad)
        except ValueError as exc:
            assert why in str(exc), exc
        else:
            raise AssertionError(f"a split asking too much was kept: {bad}")
    split = exim.check_splits(o, [{"invoice": c["key"], "amount": c["suggest"]} for c in o["candidates"]])
    chosen = {"s": {"action": "match", "context": {"kind": "settle", "remittance": rem_id, "splits": split}}}

    def paid(docs, decisions):
        got = exim.build(docs, decisions, today=today)
        return got, {r["inv_no"]: r["due"]["received_fc"] for r in got["invoices"]}

    got, by = paid(sbs + [advice], chosen)
    assert by == {"KE/26-27/101": 10000.0, "KE/26-27/102": 10000.0, "KE/26-27/103": 5000.0}, by
    assert got["remittances"][0]["available_fc"] == 0 and got["remittances"][0]["chosen"]
    assert not {"pay_confirm", "pay_unplaced"} & {r["status"] for r in got["all_records"]}

    named = sbs + [{**advice, "invoice_ref": "KE/26-27/103"}]         # the advice names the youngest
    _, by = paid(named, {})
    assert by["KE/26-27/103"] == 10000.0, "the advice's own invoice is paid first"
    o = exim.settle_options(exim.build(named, today=today), rem_id)
    assert (o["candidates"][0]["inv_no"], o["candidates"][0]["suggest"]) == ("KE/26-27/103", 10000.0)
    _, by = paid(named, chosen)
    assert by["KE/26-27/103"] == 5000.0, "the advice's own invoice overrode the user's split"


def _kotak_lines(conversion=True, colon="："):
    """A bank advice laid out like Kotak's Remittance Transaction Advice as a
    scan reads it back - title run together, full-width colons, the amount
    after its currency, a dot for a comma - with made-up parties: EUR 30,050
    from Casa Oliva, 28,000 of it repaying two packing credit loans."""
    return ["RAJPUTANA COMMERCIAL BANK LTD", "RemittanceTransactionAdvice", "Date:8-Jul-2026", "KESARI TEXTILE MILLS",
            "IN/INDIA", f"Remittance No {colon}7412XAR261900123", f"RemitterReference {colon}09MX260703000111",
            "Remitter :CASAOLIVASRL", "VIA DEI TESSITORI 18", "Credit Value Date : 8-Jul-2026",
            "ReceivedFromBank :BANCADELCHIANTISPA", "Remitter Country", "Purpose of Remittance :PO103",
            "Purpose of RemittanceDescription : Advance receiptsagainst export contracts, which will",
            "becoveredlaterbyGR/PP/SOFTEX/SDF-otherthan", "Currency Conversion details are as below :",
            "Purchase/SaleIndicatafrom CurrencyIAmountRate To CurrencyAmount",
            *(["Sales 30,050.00EUR 1.0 30,050.00EUR"] if conversion else []),
            "WeadvisehavingDebitediCreditedyouraccountas underfor charges:",
            "AccountNumber Particulars CCY DriCr Amount", "6390417001 Credit beneficiary EUR Cr 2,050.00",
            "Pre-shipmentDeal ReferenceRepaymentCCY RepaymentAmt Value date",
            "7412PCF261100111 EUR 8.000.00 04-Jul-2026", "7412PCF261100222 EUR 20,000.00 04-Jul-2026",
            'Please note that this statement/advices should not be construed as a"Tax invoices" under the Goods and '
            "Services Tax Law.", "This is a computer generated advice and does not require any signature."]


def test_export_bank_advice_read_from_a_scan():
    """A real advice of this layout was taken for an invoice (its GST
    disclaimer says "Tax invoices") and, read as an advice, would have given
    one packing credit repayment as the amount. It is an advice, the amount
    is the whole remittance, and where the money went is kept beside it."""
    from app.backend.exim_read import from_text
    got = from_text("\n".join(_kotak_lines()))
    assert got["type"] == "fira", got
    assert (got["currency"], got["amount_fc"], got["credit_date"]) == ("EUR", 30050.0, "2026-07-08"), got
    assert (got["ref"], got["purpose_code"], got["advance"]) == ("7412XAR261900123", "P0103", True), got
    assert got["credited_fc"] == 2050.0 and [r["amount_fc"] for r in got["repaid"]] == [8000.0, 20000.0], got
    assert got["remitter"] == "CASAOLIVASRL" and got["confidence"] == 0.95 and not got["notes"], got
    # Without the conversion line the parts still make the whole.
    assert from_text("\n".join(_kotak_lines(conversion=False)))["amount_fc"] == 30050.0
    # A scan that loses every space runs the rate into the amount after it, and each repayment into its date.
    glued = from_text("\n".join(ln.replace(" ", "") for ln in _kotak_lines()))
    assert (glued["amount_fc"], glued["credited_fc"]) == (30050.0, 2050.0), glued
    assert [(r["amount_fc"], r["date"]) for r in glued["repaid"]] == [(8000.0, "2026-07-04"), (20000.0, "2026-07-04")], glued


def test_export_shipping_bill_buyer_past_a_stray_letter():
    """Real ICEGATE bills put a lone letter between Part II's heading and the
    names, and one on the exporter's line in Part I: the buyer was read as
    that letter, so no remitter could ever match it."""
    from app.backend.exim_read import read_sb
    head = ["INDIAN CUSTOMS EDI SYSTEM INJAI6 7400021 13-FEB-26", "PART - I - SHIPPING BILL SUMMARY",
            "1.EXPORTER'S NAME & ADDRESS", "KESARI TEXTILE MILLS M", "PLOT 41, SITAPURA INDUSTRIAL AREA",
            "2.CONSIGNEE NAME & ADDRESS", "CASA OLIVA S.R.L", "VIA DEI TESSITORI 18",
            "PART - II - INVOICE DETAILS", "8.EXPORTER'S NAME & ADDRESS 9.BUYER'S NAME & ADDRESS", "M"]
    assert read_sb("\n".join(head + ["KESARI TEXTILE MILLS SAME AS CONSIGNEE"]))["buyer"] == "CASA OLIVA S.R.L"
    assert read_sb("\n".join(head + ["KESARI TEXTILE MILLS NORDLICHT HEIMTEXTIL GMBH"]))["buyer"] == \
        "NORDLICHT HEIMTEXTIL GMBH"


def test_export_payment_for_several_invoices_at_once():
    """One buyer's three invoices and one payment that covers them, less the
    banks' charges: offered as one question for the three, filled in the
    same way in the settle window; turned down, it is not offered again; kept,
    it pays all three - and the advance code on money for goods already
    shipped is flagged for the bank to put right."""
    from app.backend import exim
    buyer = "CASA OLIVA S.R.L"
    sbs = [_sb("KE/26-27/031", "7500031", "2026-02-10", 12400.0, "EUR", buyer),
           _sb("KE/26-27/032", "7500032", "2026-02-14", 9850.0, "EUR", buyer),
           _sb("KE/26-27/033", "7500033", "2026-02-19", 14230.0, "EUR", buyer),
           _sb("KE/26-27/030", "7500030", "2026-01-05", 5000.0, "EUR", "NORDLICHT HEIMTEXTIL GMBH")]
    advice = {**_advice(36420.0, "2026-07-08", remitter="CASAOLIVASRL", currency="EUR"),
              "purpose_code": "P0103", "advance": True}
    today = D("2026-09-27")
    view = exim.build(sbs + [advice], today=today)
    ask = next(r for r in view["all_records"] if r["status"] == "pay_confirm")
    assert ask["counterpart"]["kind"] == "set" and "KE/26-27/031, KE/26-27/032 and KE/26-27/033" in ask["why"]["title"]
    split = [(s["invoice"], s["amount"]) for s in ask["counterpart"]["splits"]]
    keys = {r["inv_no"]: r["key"] for r in view["invoices"]}
    assert split == [(keys["KE/26-27/031"], 12400.0), (keys["KE/26-27/032"], 9850.0), (keys["KE/26-27/033"], 14170.0)]
    assert any("P0103" in c for c in ask["why"]["checks"])
    o = exim.settle_options(view, ask["counterpart"]["remittance"])
    assert [(c["inv_no"], c["suggest"]) for c in o["candidates"]][:3] == [
        ("KE/26-27/031", 12400.0), ("KE/26-27/032", 9850.0), ("KE/26-27/033", 14170.0)], o["candidates"]

    refused = exim.build(sbs + [advice], {ask["fp"]: {"action": "ignore", "context": ask["counterpart"]}}, today=today)
    assert all(r["counterpart"]["kind"] != "set" for r in refused["all_records"] if r["status"] == "pay_confirm")

    chosen = {"s": {"action": "match", "context": {"kind": "settle", "remittance": ask["counterpart"]["remittance"],
                                                  "splits": [{"invoice": k, "amount": a} for k, a in split]}}}
    got = exim.build(sbs + [advice], chosen, today=today)
    due = {r["inv_no"]: r["due"] for r in got["invoices"]}
    assert all(due[n]["state"] == "paid" for n in ("KE/26-27/031", "KE/26-27/032", "KE/26-27/033"))
    assert due["KE/26-27/033"]["short_by_charges"] == 60.0
    flag = next(r for r in got["all_records"] if r["status"] == "advance_shipped")
    assert "7500031" in flag["why"]["todo"] and "P0102" in flag["why"]["todo"], flag["why"]

    # Two runs that both fit are a guess: no set is offered.
    even = [_sb(f"KE/26-27/04{i}", f"750004{i}", f"2026-03-0{i}", 10000.0, "EUR", buyer) for i in range(1, 5)]
    view = exim.build(even + [{**advice, "amount_fc": 20000.0}], today=today)
    assert all(r["counterpart"]["kind"] != "set" for r in view["all_records"] if r["status"] == "pay_confirm")


def test_export_document_read_again():
    """A file read before the reader knew its layout is read again from the
    kept copy: the new reading replaces the old, a correction that still says
    something else stays, one for another kind of document goes, and the
    document waits to be checked again. A second read that finds nothing new
    changes nothing."""
    import asyncio
    import io

    from reportlab.lib.pagesizes import A4
    from reportlab.pdfgen import canvas

    from app.backend import server, store
    buf = io.BytesIO()
    c = canvas.Canvas(buf, pagesize=A4)
    for n, line in enumerate(_kotak_lines(colon=":")):
        c.drawString(40, 800 - 14 * n, line)
    c.save()
    ws = store.create_workspace(name="Read again test", module="exports")
    try:
        doc_id, _ = store.add_document(ws, "advice scan.pdf", buf.getvalue(), {"type": "invoice", "inv_no": None},
                                       folder="Invoices", name="advice scan.pdf", status="posted")
        store.save_corrections(ws, doc_id, {"remitter": "CASA OLIVA SRL", "inv_no": "KE/26-27/031"})
        asyncio.run(server.exports_doc_reread(ws, doc_id))
        doc = store.list_documents(ws)[0]
        assert (doc["type"], doc["amount_fc"], doc["status"]) == ("fira", 30050.0, "pending"), doc
        assert doc["corrections"] == {"remitter": "CASA OLIVA SRL"}, doc["corrections"]
        assert doc["kept_as"] == "Bank advice 7412XAR261900123 of 08-Jul-2026.pdf" and doc["folder"] == "Bank advices"
        log = store.list_audit(ws)
        assert log[0]["what"] == "reread" and log[0]["changes"][0]["field"] == "Kind of document", log[0]
        asyncio.run(server.exports_doc_reread(ws, doc_id))
        assert "reads the same" in store.list_audit(ws)[0]["detail"]
        assert store.list_documents(ws)[0]["kept_as"] == doc["kept_as"]
    finally:
        store.delete_workspace(ws)


def test_export_look_up_on_icegate():
    """A refund with no status on file carries what ICEGATE's shipping bill
    enquiry asks for - port, number, and the date the way it wants it - and a
    link to the page, on the invoice, on its open entry and in the look-up
    sheet. The app itself never goes to ICEGATE."""
    import io

    import openpyxl

    from app.backend import exim, exim_drafts
    sb = _sb("KE/26-27/071", "7600071", "2026-02-12", 8000.0)
    sb["invoices"][0]["igst_amount"] = 14000.0
    register = {"id": "o1", "name": "register.xlsx", "type": "own_register", "rows": [
        {"sheet": "2025-26", "row": 9, "sb_no": "7600071", "inv_no": "KE/26-27/071", "igst": 14000.0}]}
    view = exim.build([sb, register], today=D("2026-09-27"))
    row = view["invoices"][0]
    assert (row["icegate"]["port"], row["icegate"]["sb_no"], row["icegate"]["sb_date"]) == ("INJAI6", "7600071", "12-02-2026")
    assert row["icegate"]["sb_enquiry"].startswith("https://foservices.icegate.gov.in/")
    entry = next(r for r in view["all_records"] if r["status"] == "igst_pending")
    assert entry["icegate"]["sb_no"] == "7600071", "the open entry does not say where to look it up"
    ws = openpyxl.load_workbook(io.BytesIO(exim_drafts.lookup_workbook(view)))["Look up on ICEGATE"]
    assert ws.cell(4, 1).value == "7600071" and ws.cell(4, 2).number_format == "dd-mm-yyyy"
    assert ws.cell(4, 11).hyperlink.target == row["icegate"]["sb_enquiry"] and "captcha" in ws.cell(2, 1).value


def test_export_lists_picked_up_from_downloads():
    """An ICEGATE status download saved in the last week is offered; a file
    that is not an export list, an old one, one already in the register and
    one turned down are not. Taking it files it to be checked - and the
    Downloads folder is left exactly as it was."""
    import asyncio
    import os
    import tempfile
    import time
    from pathlib import Path

    import openpyxl

    from app.backend import server, store

    folder = Path(tempfile.mkdtemp(prefix="downloads-"))

    def sheet(name, header, rows, age_days=0):
        wb = openpyxl.Workbook()
        wb.active.append(header)
        for r in rows:
            wb.active.append(r)
        wb.save(folder / name)
        if age_days:
            old = time.time() - age_days * 86400
            os.utime(folder / name, (old, old))

    status = ["Shipping Bill No", "Shipping Bill Date", "Port Code", "Response Code", "Status", "IGST Amount"]
    sheet("IGST_refund_status.xlsx", status, [["7600071", "12-02-2026", "INJAI6", "SB005", "Invoice mismatch", 14000]])
    sheet("IGST_refund_status (1).xlsx", status, [["7600072", "14-02-2026", "INJAI6", "SB000", "Validated", 9000]])
    sheet("old_status.xlsx", status, [["7600060", "01-01-2026", "INJAI6", "SB002", "EGM pending", 5000]], age_days=10)
    sheet("price list.xlsx", ["Item", "Rate"], [["Cushion cover", 4.5]])
    before = {p.name: p.read_bytes() for p in folder.iterdir()}
    was, os.environ["ASSURANCE_DOWNLOADS_DIR"] = os.environ.get("ASSURANCE_DOWNLOADS_DIR"), str(folder)
    ws = store.create_workspace(name="Downloads test", module="exports")
    try:
        found = asyncio.run(server.exports_downloads(ws))["files"]
        assert sorted(f["name"] for f in found) == ["IGST_refund_status (1).xlsx", "IGST_refund_status.xlsx"], found
        assert {f["type"] for f in found} == {"refund_status"} and "path" not in found[0]
        take = next(f for f in found if f["name"] == "IGST_refund_status.xlsx")
        view = asyncio.run(server.exports_downloads_take(ws, {"add": take["id"]}))
        assert [d["type"] for d in view["pending"]] == ["refund_status"], "it did not wait to be checked"
        assert "from your Downloads folder" in store.list_audit(ws)[0]["detail"]
        other = next(f for f in found if f["name"] != "IGST_refund_status.xlsx")
        left = asyncio.run(server.exports_downloads_take(ws, {"skip": other["id"]}))["files"]
        assert left == [], left
        try:
            asyncio.run(server.exports_downloads_take(ws, {"add": "0" * 64}))
        except Exception as exc:                          # noqa: BLE001 - HTTPException
            assert getattr(exc, "status_code", None) == 404
        else:
            raise AssertionError("a file not found in Downloads was taken")
        assert {p.name: p.read_bytes() for p in folder.iterdir()} == before, "the Downloads folder was changed"
    finally:
        store.save_settings({"downloads_skipped": []})
        store.delete_workspace(ws)
        if was is None:
            os.environ.pop("ASSURANCE_DOWNLOADS_DIR", None)
        else:
            os.environ["ASSURANCE_DOWNLOADS_DIR"] = was


def test_export_your_own_yes_or_no():
    """Where no document says whether the return is filed or the refunds came
    back, the user's yes or no stands in; and it wins over a document the user
    knows is wrong."""
    from app.backend import exim
    sb = _sb("KE/26-27/111", "7300001", "2026-04-10", 8000.0)
    sb["invoices"][0]["igst_amount"] = 12000.0                      # WPAY: a refund to come back
    sb["dbk_amount"] = 3000.0
    april = {"id": "g1", "name": "GSTR-1.json", "type": "gstr1", "gstin": "08AMHCK7032W1ZX",
             "period": "2026-04", "lines": []}                        # April's return, without it
    today = D("2026-09-20")
    key = exim.build([sb], today=today)["invoices"][0]["key"]

    def said(**answers):
        return {f"m-{what}": {"action": "mark", "updated_at": "2026-09-20T10:00:00",
                              "context": {"kind": "mark", "invoice": key, "what": what, "value": v}}
                for what, v in answers.items()}

    got = exim.build([sb, april], today=today)
    row = got["invoices"][0]
    assert "not_in_gst" in {r["status"] for r in got["all_records"]}
    assert (row["refund"]["igst"]["state"], row["marks"]["igst"]["value"]) == ("unknown", None)

    got = exim.build([sb, april], said(gstr1="yes", igst="yes", dbk="no"), today=today)
    row, recs = got["invoices"][0], {r["status"]: r for r in got["all_records"]}
    assert not {"not_in_gst", "igst_pending"} & set(recs), set(recs)
    assert row["chips"]["gst"] == {"state": "ok", "text": "Filed · marked by you"}
    assert (row["refund"]["igst"]["state"], row["refund"]["igst"]["recd"]) == ("received", 12000.0)
    assert "You marked it as not received." in recs["dbk_pending"]["why"]["what"]
    assert row["marks"]["gstr1"] == {"applies": True, "doc": "no", "said": "yes",
                                     "at": "2026-09-20T10:00:00", "value": "yes"}

    got = exim.build([sb], said(gstr1="no"), today=today)             # no return added at all
    rec = next(r for r in got["all_records"] if r["status"] == "not_in_gst")
    assert rec["tone"] == "serious" and "You marked it as not filed" in rec["why"]["what"]

    paid = {"id": "rs", "name": "refunds.xlsx", "type": "refund_status", "rows": [
        {"sb_no": "7300001", "kind": "igst", "scroll_no": "S-1", "scroll_date": "2026-06-01", "amount": 12000.0}]}
    assert exim.build([sb, paid], today=today)["invoices"][0]["refund"]["igst"]["state"] == "received"
    got = exim.build([sb, paid], said(igst="no"), today=today)
    ig = got["invoices"][0]["refund"]["igst"]
    assert (ig["state"], ig["doc_state"]) == ("pending", "received"), ig
    assert got["invoices"][0]["marks"]["igst"] == {"applies": True, "doc": "yes", "said": "no",
                                                   "at": "2026-09-20T10:00:00", "value": "no"}


def test_export_corrections_laid_over_what_was_read():
    """A bill of lading read with the wrong invoice number is filed against
    the wrong invoice until the number is corrected; then it is filed right.
    What was read stays as it was, each change is told old-to-new, a value put
    back is no longer a correction, and a bad value is refused whole."""
    from app.backend import exim, exim_fields as xf
    sbs = [_sb("KE/26-27/013", "7400013", "2026-05-02", 5000.0), _sb("KE/26-27/014", "7400014", "2026-05-03", 6000.0)]
    bl = {"id": "b1", "name": "bl.pdf", "type": "bl", "bl_no": "SLL260500017", "sob_date": "2026-05-06",
          "invoice_refs": [{"inv_no": "KE/26-27/013", "inv_date": "2026-05-03"}], "sb_refs": [], "containers": []}
    today = D("2026-06-01")

    def filed(doc):
        return {r["inv_no"]: [b["bl_no"] for b in r["bls"]] for r in exim.build(sbs + [doc], today=today)["invoices"]}

    assert filed(bl) == {"KE/26-27/013": ["SLL260500017"], "KE/26-27/014": []}
    fixes, changes = xf.correct(bl, {}, {"fields": {"invoice_refs": "KE/26-27/014", "bl_no": "SLL260500017 "}})
    assert changes == [{"field": "Invoice numbers on it", "old": "KE/26-27/013", "new": "KE/26-27/014"}], changes
    assert bl["invoice_refs"][0]["inv_no"] == "KE/26-27/013", "what was read was changed"
    assert filed(xf.apply(bl, fixes)) == {"KE/26-27/013": [], "KE/26-27/014": ["SLL260500017"]}
    back, _ = xf.correct(bl, fixes, {"fields": {"invoice_refs": "KE/26-27/013"}})
    assert back == {} and xf.apply(bl, back)["invoice_refs"][0]["inv_date"] == "2026-05-03", back
    for edit, why in (({"fields": {"sob_date": "31/02/2026"}}, "is not a date"),
                      ({"fields": {"weight": "12"}}, "has no field"),
                      ({"type": "gstr1"}, "cannot be turned")):
        try:
            xf.correct(bl, fixes, edit)
        except ValueError as exc:
            assert why in str(exc), exc
        else:
            raise AssertionError(f"a bad edit was kept: {edit}")

    sb = {**sbs[0], "invoices": [{**sbs[0]["invoices"][0], "igst_amount": 0.0}], "igst_amount": 0.0, "scheme": "LUT"}
    fixes, changes = xf.correct(sb, {}, {"tables": {"invoices": [{"_i": 0, "igst_amount": "9,000"},
                                                                 {"inv_no": "KE/26-27/015", "value_fc": "700"}]}})
    now = xf.apply(sb, fixes)
    assert (now["igst_amount"], now["scheme"], len(now["invoices"])) == (9000.0, "WPAY", 2), now
    assert [c["field"] for c in changes] == ["Invoices on it, KE/26-27/013: IGST (₹)", "Invoices on it: row added"]
    for rows, why in (([{"_i": 0}, {"value_fc": "5"}], "has no invoice number"),
                      ([{"_i": 0, "value_fc": "-4"}], "zero or more")):
        try:
            xf.correct(sb, {}, {"tables": {"invoices": rows}})
        except ValueError as exc:
            assert why in str(exc), exc
        else:
            raise AssertionError(f"a bad row was kept: {rows}")
    register = {"type": "own_register", "rows": [{"sheet": "2026-27", "row": 4, "sb_no": "7400013", "inv_no": "KE/26-27/013"}]}
    try:
        xf.correct(register, {}, {"tables": {"rows": [{"_i": 0}, {"sb_no": "7400099"}]}})
    except ValueError as exc:
        assert "add them to the file itself" in str(exc)
    else:
        raise AssertionError("a row was added to the user's own register")

    unknown = {"type": "unknown", "notes": ["Could not tell what this document is from its text."]}
    assert [k for k, _ in xf.form(unknown, {})["types"]] == ["sb", "bl", "fira", "invoice"]
    fixes, changes = xf.correct(unknown, {}, {"type": "bl", "fields": {"bl_no": "HLCU123456", "invoice_refs": "KE/26-27/014"}})
    assert changes[0] == {"field": "Kind of document", "old": "Not recognised", "new": "Bill of lading"}
    assert filed({**xf.apply(unknown, fixes), "id": "u1", "name": "scan.jpg"})["KE/26-27/014"] == ["HLCU123456"]
    assert [f for f, _ in xf.facts(xf.apply(bl, {}))] == ["Bill of lading", "Shipped on board", "Invoice numbers"]


def test_export_documents_wait_to_be_checked():
    """A document added for checking counts nowhere until it is posted; the
    corrections are kept beside what was read; the kept file takes its
    corrected name; and the history keeps each step, newest first."""
    from app.backend import store
    ws = store.create_workspace(name="Check test", module="exports")
    meta = {"type": "bl", "bl_no": "SLL1", "invoice_refs": [{"inv_no": "KE/26-27/013", "inv_date": None}]}
    doc_id, fresh = store.add_document(ws, "scan.pdf", b"%PDF-1.4 a scanned bill of lading", meta,
                                       folder="Bills of lading", name="BL SLL1 for KE-26-27-013.pdf",
                                       group="KESARI TEXTILE MILLS - exports", status="pending")
    try:
        assert fresh and store.list_documents(ws, "posted") == []
        assert [d["status"] for d in store.list_documents(ws)] == ["pending"]
        store.save_corrections(ws, doc_id, {"invoice_refs": [{"inv_no": "KE/26-27/014", "inv_date": None}]})
        doc = store.list_documents(ws, "pending")[0]
        assert doc["invoice_refs"][0]["inv_no"] == "KE/26-27/013" and \
            doc["corrections"]["invoice_refs"][0]["inv_no"] == "KE/26-27/014", "what was read was overwritten"
        path = store.refile_document(ws, doc_id, "Bills of lading", "BL SLL1 for KE-26-27-014.pdf")
        assert path and path.endswith("BL SLL1 for KE-26-27-014.pdf") and store.get_document(ws, doc_id)["path"].exists()
        assert store.post_documents(ws, [doc_id, "not-a-document"]) == [doc_id]
        assert store.post_documents(ws, [doc_id]) == [], "a posted document was posted twice"
        assert [d["id"] for d in store.list_documents(ws, "posted")] == [doc_id]
        store.audit(ws, "added", "Added a bill of lading.", doc_id)
        store.audit(ws, "corrected", "Corrected it.", doc_id,
                    [{"field": "Invoice numbers on it", "old": "KE/26-27/013", "new": "KE/26-27/014"}], by="Test")
        log = store.list_audit(ws)
        assert [a["what"] for a in log] == ["corrected", "added"] and log[0]["changes"][0]["new"] == "KE/26-27/014"
    finally:
        store.delete_workspace(ws)
    assert store.list_audit(ws) == [], "a deleted register left its history behind"


def test_database_from_before_the_check_gains_it():
    """A database made before documents waited to be checked opens with every
    document it holds posted, and nothing else changed."""
    import sqlite3
    import tempfile
    from pathlib import Path

    from app.backend import store
    conn = sqlite3.connect(Path(tempfile.mkdtemp(prefix="olddb-")) / "reconcile.db")
    conn.row_factory = sqlite3.Row
    conn.executescript("""
        CREATE TABLE runs (id TEXT PRIMARY KEY, workspace_id TEXT, created_at TEXT NOT NULL);
        CREATE TABLE workspaces (id TEXT PRIMARY KEY, created_at TEXT, updated_at TEXT, module TEXT);
        CREATE TABLE documents (id TEXT PRIMARY KEY, workspace_id TEXT NOT NULL, added_at TEXT NOT NULL,
            filename TEXT, sha256 TEXT, type TEXT, size INTEGER, meta_json TEXT, stored_path TEXT);
        INSERT INTO documents VALUES ('d1', 'w1', '2026-09-01T10:00:00', 'sb.pdf', 'x', 'sb', 10, '{}', 'a/b/sb.pdf');
        PRAGMA user_version = 4;""")
    store._migrate(conn)
    row = conn.execute("SELECT status, corrections, meta_json, stored_path FROM documents").fetchone()
    assert (row["status"], row["corrections"], row["meta_json"], row["stored_path"]) == ("posted", None, "{}", "a/b/sb.pdf")
    conn.close()


def test_export_invoice_numbers_that_differ():
    from app.backend.exim import _sb005_route
    from app.backend.exim_read import ocr_key
    assert ocr_key("VI-0484") == ocr_key("V1-0484") == ocr_key("VL-0484")   # OCR slips link
    typo, fix = _sb005_route({"invoice_doc": None}, {"inum": "KE/25-26/033"}, {"inv_no": "KF/25-26/033"})
    assert fix == "check" and "typing slip" in typo          # which side is wrong needs the invoice
    series, fix = _sb005_route({"invoice_doc": None}, {"inum": "KE/25-26/047"}, {"inv_no": "CI/2611"})
    assert fix == "port" and "second series" in series
    _, fix = _sb005_route({"invoice_doc": {"inv_no": "KF/25-26/033"}}, {"inum": "KE/25-26/033"},
                          {"inv_no": "KF/25-26/033"})
    assert fix == "9a"                                        # the invoice sides with the shipping bill
    _, fix = _sb005_route({"invoice_doc": {"inv_no": "KE/25-26/033"}}, {"inum": "KE/25-26/033"},
                          {"inv_no": "KF/25-26/033"})
    assert fix == "port"                                      # the invoice sides with Table 6A
    dated, fix = _sb005_route({"invoice_doc": None}, {"inum": "KE/25-26/066"}, {"inv_no": "KE/25-26/066"})
    assert fix == "check" and "only the date" in dated


def test_export_sample_finds_every_planted_problem():
    """The whole sample, then every draft made from it: each file opens, and
    the SB005 challans come to Rs 1,000 a shipping bill."""
    import io
    import zipfile
    from xml.dom import minidom

    import openpyxl

    from app.backend import exim_drafts
    from tools.synth_exim import verify
    view = verify(quiet=True)
    who = {"name": "KESARI TEXTILE MILLS", "gstin": "08AMHCK7032W1ZX"}
    today = D("2026-09-26")

    def opens(name, data):
        if name.endswith(".docx"):
            minidom.parseString(zipfile.ZipFile(io.BytesIO(data)).read("word/document.xml"))
        elif name.endswith(".xlsx"):
            return openpyxl.load_workbook(io.BytesIO(data))

    pack = zipfile.ZipFile(io.BytesIO(exim_drafts.sb005_pack(view, who, today)))
    for name in pack.namelist():
        opens(name, pack.read(name))
    challans = opens("x.xlsx", pack.read("Challan schedule.xlsx"))["Challan schedule"]
    fees = [r[3] for r in challans.iter_rows(min_row=4, values_only=True) if isinstance(r[3], (int, float))]
    assert sum(fees) == 1000 * exim_drafts.summary(view)["sb005"]["count"], fees
    for made, name in ((exim_drafts.table9a_workbook(view), "x.xlsx"), (exim_drafts.lookup_workbook(view), "x.xlsx"),
                       (exim_drafts.own_format_workbook(view), "x.xlsx"),
                       (exim_drafts.sb003_undertaking(view, who, today), "x.docx")):
        opens(name, made)
    for made in (exim_drafts.declaration(view, who, today), exim_drafts.extension(view, who, today)):
        z = zipfile.ZipFile(io.BytesIO(made))
        for name in z.namelist():
            opens(name, z.read(name))


def test_export_hand_typed_dates():
    from app.backend.exim_read import _date
    assert _date("12-04-25.") == _date("12.04.25") == _date("12/4/25") == "2025-04-12"


def test_documents_are_kept_by_what_they_are():
    from app.backend.exim_read import kept_name
    from app.backend.store import file_name
    sb = {"type": "sb", "sb_no": "3183667", "sb_date": "2024-08-20", "port": "INNSA1",
          "invoices": [{"inv_no": "KE/24-25/088"}]}
    raw = "318366720082024INNSA1SB2200820241105.PDF"
    assert file_name(raw, kept_name(sb, raw)) == "SB 3183667 of 20-Aug-2024 (INNSA1) - KE-24-25-088.pdf"
    assert kept_name({"type": "gstr1", "period": "2026-01"}, "r.json") == "GSTR-1 Jan-2026.json"
    assert kept_name({"type": "edpms"}, "EDPMS statement.xlsx") == "EDPMS statement.xlsx"   # a sheet keeps its name
    assert file_name("Book1.xlsx", "Tally - Book1.xlsx") == "Tally - Book1.xlsx"
    assert file_name('a:b?"c.XLSX') == "a b c.xlsx"
    assert file_name("", None) == "document"


def test_ocr_fallback_only_for_scans_and_broken_text():
    """A PDF that is simply not an export document must not be read again by
    OCR (that cost 6 s a file on real downloads); a scan with a text layer on
    top, or text that is not words, must be."""
    import io as _io
    from PIL import Image as _Image
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.utils import ImageReader
    from reportlab.pdfgen import canvas
    from app.backend.exim_read import _text_layer_unreliable

    def pdf(text, picture=False):
        buf = _io.BytesIO()
        c = canvas.Canvas(buf, pagesize=A4)
        if picture:
            c.drawImage(ImageReader(_Image.new("RGB", (60, 85), "white")), 0, 0, *A4)
        c.drawString(72, 720, text)
        c.save()
        return buf.getvalue(), text

    plain = "Quarterly maintenance agreement between the parties for office equipment and services"
    assert not _text_layer_unreliable(*pdf(plain))
    assert _text_layer_unreliable(*pdf(plain, picture=True))              # a scan with text laid over it
    assert _text_layer_unreliable(*pdf("(cid:12)(cid:7)(cid:30) (cid:4)(cid:9)"))
    assert _text_layer_unreliable(*pdf("7XQ 4/9 8#2 %%1 0Z9 K7P 22Q"))    # not words



def _stand_in(handle):
    """A web server on this computer standing in for one on the internet:
    handle(method, path, headers, body) gives (status, headers, body - bytes, or
    lines sent one by one). Returns the server and its address."""
    import http.server
    import threading

    class Handler(http.server.BaseHTTPRequestHandler):
        def _answer(self):
            n = int(self.headers.get("Content-Length") or 0)
            status, headers, body = handle(self.command, self.path, self.headers, self.rfile.read(n) if n else b"")
            self.send_response(status)
            if isinstance(body, bytes):
                headers = {**headers, "Content-Length": str(len(body))}
            for k, v in headers.items():
                self.send_header(k, v)
            self.end_headers()
            for chunk in [body] if isinstance(body, bytes) else body:
                self.wfile.write(chunk)
                self.wfile.flush()

        do_GET = do_POST = _answer

        def log_message(self, *args):
            pass

    server = http.server.ThreadingHTTPServer(("127.0.0.1", 0), Handler)
    threading.Thread(target=server.serve_forever, daemon=True).start()
    return server, f"http://127.0.0.1:{server.server_address[1]}"


def test_ai_checker_fetched_checked_and_removed():
    """The AI checker's runtime comes from the release page, carries on from a
    download cut off, is checked against the published checksum (a mismatch is
    deleted, never unpacked), loses the old CUDA 12 libraries, and the model is
    pulled through the runtime. Removing it leaves nothing behind."""
    import hashlib
    import io
    import json
    import tempfile
    import zipfile
    from pathlib import Path

    from app.backend import ai
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w") as z:
        for name in ("ollama.exe", "lib/ollama/ggml-base.dll", "lib/ollama/cuda_v12/ggml-cuda.dll",
                     "lib/ollama/cuda_v13/ggml-cuda.dll"):
            z.writestr(name, f"stand-in for {name}" * 200)
    pack = buf.getvalue()
    sums = {"sum": hashlib.sha256(pack).hexdigest()}
    asked = []

    def handle(method, path, headers, body):
        asked.append((path, headers.get("Range")))
        if path.endswith("/sha256sum.txt"):
            return 200, {}, f"{'0' * 64}  ./ollama-linux-amd64.tgz\n{sums['sum']}  ./ollama-windows-amd64.zip\n".encode()
        if path.endswith("/ollama-windows-amd64.zip"):
            start = int(headers["Range"].split("=")[1].rstrip("-")) if headers.get("Range") else 0
            return (206 if start else 200), {}, pack[start:]
        if path == "/api/pull":
            assert json.loads(body)["model"] == ai.MODEL
            lines = [{"status": "pulling manifest"}, {"digest": "sha256:aa", "total": 3000, "completed": 1000},
                     {"digest": "sha256:aa", "total": 3000, "completed": 3000}, {"status": "success"}]
            return 200, {"Content-Type": "application/x-ndjson"}, [json.dumps(x).encode() + b"\n" for x in lines]
        return 404, {}, b""

    server, url = _stand_in(handle)
    before = (ai.FOLDER, ai.RELEASES, ai.GIVEN, ai.SHARED)
    try:
        with tempfile.TemporaryDirectory() as tmp:
            ai.FOLDER, ai.RELEASES, ai.GIVEN, ai.SHARED = Path(tmp) / "ai", url, "", Path(tmp) / "Shared AI"
            part = ai.FOLDER / "downloads" / f"ollama-{ai.OLLAMA_VERSION}-windows-amd64.zip.part"
            part.parent.mkdir(parents=True)
            part.write_bytes(pack[:5000])                     # a download cut off part-way
            ai._get_runtime()
            assert (f"/{ai.OLLAMA_VERSION}/ollama-windows-amd64.zip", "bytes=5000-") in asked, asked
            lib = ai.runtime_dir() / "lib" / "ollama"
            assert (ai.runtime_dir() / "ollama.exe").exists() and (lib / "cuda_v13").exists()
            assert not (lib / "cuda_v12").exists(), "the CUDA 12 libraries were kept"
            assert not (ai.FOLDER / "downloads").exists() and not ai.installed()     # no model yet

            ai.GIVEN = url                                   # the model comes through the runtime's own pull
            ai._pull()
            assert (ai.models_dir() / "MODEL").read_text() == ai.MODEL
            ai.GIVEN = ""
            assert ai.installed()
            ai.remove()
            assert not ai.FOLDER.exists() and ai.status()["status"] == "absent"

            sums["sum"] = "f" * 64                           # a runtime that is not the one published
            try:
                ai._get_runtime()
                raise AssertionError("a runtime that failed its checksum was unpacked")
            except ai.AiError as exc:
                assert "checksum" in str(exc)
            assert not ai.runtime_dir().exists()
            assert not list((ai.FOLDER / "downloads").glob("*.zip")), "the mismatched download was kept"
    finally:
        ai.FOLDER, ai.RELEASES, ai.GIVEN, ai.SHARED = before
        ai._set(status="absent", done=0, total=0, step="", message="")
        server.shutdown()


def test_ai_checker_reads_documents_again_in_the_background():
    """A posted bank advice and a scanned bill of lading are read again by the
    model (the advice as its text, the scan as a picture), under a fixed
    instruction and an answer held to their fields. Where it reads the advice's
    amount differently the register lists it; using its reading keeps it as a
    correction in the history, and the advice is checked again and agrees. A
    field the model leaves blank is not counted against the app."""
    import asyncio
    import io
    import json

    from fastapi import HTTPException
    from PIL import Image
    from reportlab.lib.pagesizes import A4
    from reportlab.pdfgen import canvas

    from app.backend import ai, ai_check, server, store
    buf = io.BytesIO()
    c = canvas.Canvas(buf, pagesize=A4)
    for n, line in enumerate(_kotak_lines(colon=":")):
        c.drawString(40, 800 - 14 * n, line)
    c.save()
    picture = io.BytesIO()
    Image.new("RGB", (400, 560), "white").save(picture, "PNG")
    advice = {"type": "fira", "ref": "7412XAR261900123", "credit_date": "2026-07-08", "currency": "EUR",
              "amount_fc": 2050.0, "remitter": "CASA OLIVA SRL", "purpose_code": "P0103", "advance": True,
              "invoice_ref": None}
    bl = {"type": "bl", "bl_no": "KESBL26070011", "sob_date": "2026-07-01",
          "invoice_refs": [{"inv_no": "KE/26-27/031", "inv_date": None}], "sb_refs": [], "containers": [],
          "read_by": "recognition"}
    said = {"fira": {"kind": "fira", "amount": "30,050.00EUR", "credit_date": "8-Jul-2026", "ref": "7412XAR261900123",
                     "remitter": "CASAOLIVASRL", "purpose_code": "PO103"},
            "bl": {"kind": "bl", "bl_no": "KESBL26070011", "sob_date": "",     # a number copied with its date
                   "inv_nos": ["KE/26-27/031-04/07/2026"], "sb_nos": []}}
    asked = []

    def handle(method, path, headers, body):
        if path == "/api/version":
            return 200, {}, b'{"version": "0.34.4"}'
        if path == "/api/show":
            return 200, {}, json.dumps({"capabilities": ["completion", "vision", "thinking"]}).encode()
        if path == "/api/chat":
            ask = json.loads(body)
            asked.append(ask)
            kind = "bl" if "bl_no" in ask["format"]["properties"] else "fira"
            return 200, {}, json.dumps({"message": {"content": json.dumps(said[kind])}}).encode()
        return 404, {}, b""

    stand_in, url = _stand_in(handle)
    before = ai.GIVEN
    ws = store.create_workspace(name="AI checker test", module="exports")
    try:
        ai.GIVEN = url
        fira_id, _ = store.add_document(ws, "advice.pdf", buf.getvalue(), advice, folder="Bank advices",
                                        name="advice.pdf", status="posted")
        bl_id, _ = store.add_document(ws, "bl scan.png", picture.getvalue(), bl, folder="Bills of lading",
                                      name="bl scan.png", status="posted")
        docs = {d["id"]: d for d in server._exim_docs(ws)}
        assert docs[fira_id]["ai"] == {"state": "waiting"}, docs[fira_id]["ai"]
        assert ai_check.run_once() is not None and ai_check.run_once() is not None
        assert ai_check.run_once() is None, "a document was checked twice with nothing changed"

        by_kind = {("bl_no" in a["format"]["properties"]): a for a in asked}
        text_ask, scan_ask = by_kind[False], by_kind[True]
        assert text_ask["messages"][0]["content"] == ai_check.SYSTEM and text_ask["think"] is False
        assert text_ask["options"]["temperature"] == 0 and "7412XAR261900123" in text_ask["messages"][1]["content"]
        assert not text_ask["messages"][1].get("images"), "an advice with good text was sent as a picture"
        assert len(scan_ask["messages"][1]["images"]) == 1 and "attached pictures" in scan_ask["messages"][1]["content"]
        assert set(text_ask["format"]["required"]) == {"kind", "amount", "credit_date", "ref", "remitter", "purpose_code"}

        docs = {d["id"]: d for d in server._exim_docs(ws)}
        a = docs[fira_id]["ai"]
        assert a["state"] == "differs" and [d["label"] for d in a["differences"]] == ["Amount"], a
        assert a["differences"][0]["ai"] == "EUR 30,050.00" and a["differences"][0]["app"] == "EUR 2,050.00"
        b = docs[bl_id]["ai"]
        assert b["state"] == "agrees" and b["unread"] == ["Shipped on board", "Shipping bill numbers on it"], b
        entry = next(r for r in server._register_view(ws)["all_records"] if r["status"] == "ai_differs")
        assert entry["doc_id"] == fira_id and entry["why"]["title"].endswith("differently: amount."), entry["why"]

        try:
            asyncio.run(server.exports_doc_ai_use(ws, bl_id))
            raise AssertionError("the AI's reading was used where it agrees")
        except HTTPException as exc:
            assert exc.status_code == 400
        view = asyncio.run(server.exports_doc_ai_use(ws, fira_id))
        assert view["changes"] == 1 and not any(r["status"] == "ai_differs" for r in view["all_records"])
        doc = next(d for d in store.list_documents(ws) if d["id"] == fira_id)
        assert doc["corrections"] == {"amount_fc": 30050.0}, doc["corrections"]
        log = store.list_audit(ws)[0]
        assert log["detail"].startswith("Used the AI checker's reading of"), log
        assert log["changes"] == [{"field": "Amount", "old": "2,050.00", "new": "30,050.00"}], log["changes"]
        assert {d["id"]: d for d in server._exim_docs(ws)}[fira_id]["ai"]["state"] == "waiting"

        def stopped(*args, **kwargs):                         # the bridge's installer stopped its runtime
            raise ai.Stopped("The AI runtime stopped while it was reading.")
        real_chat, ai.chat = ai.chat, stopped
        try:
            assert ai_check.run_once() is None
        finally:
            ai.chat = real_chat
        assert {d["id"]: d for d in server._exim_docs(ws)}[fira_id]["ai"]["state"] == "waiting", \
            "a check the runtime was stopped in was kept as failed"
        ai_check.run_once()                                   # read again as it now stands
        assert {d["id"]: d for d in server._exim_docs(ws)}[fira_id]["ai"]["state"] == "agrees"
    finally:
        ai.GIVEN = before
        store.delete_workspace(ws)
        stand_in.shutdown()


def test_ai_checker_shares_one_copy_with_the_bridge():
    """One copy of the AI serves both apps. With none of its own, the checker runs
    the one Invoice Bridge lists in the shared folder - as a process of its own,
    with Ollama's pruning off, writing nothing there - and a runtime stopped part-way
    is started again, not counted as a failure. Its own complete copy is listed in
    ASCII JSON and run instead; removing it takes the note away first, then stops
    what runs from it - only that - and the checker goes on with the bridge's."""
    import json
    import os
    import shutil
    import subprocess
    import tempfile
    import types
    from pathlib import Path

    from app.backend import ai
    runtimes, stop_it = [], {"now": False}

    class Runtime:
        """Stands in for `ollama.exe serve`: what it was started with, and whether it still runs."""
        def __init__(self, args, env=None, **kw):
            self.args, self.env, self.kw, self.code, self._handle = args, env, kw, None, 0
            runtimes.append(self)

        def poll(self):
            return self.code

        def wait(self, timeout=None):
            if self.code is None:
                raise subprocess.TimeoutExpired(self.args, timeout)
            return self.code

        def kill(self):
            self.code = -9

    def handle(method, path, headers, body):
        if path == "/api/version":
            return 200, {}, b'{"version": "0.34.4"}'
        if path == "/api/show":
            return 200, {}, b'{"capabilities": ["completion"]}'
        if path == "/api/chat" and stop_it["now"]:
            runtimes[-1].code = 1                          # stopped part-way: the answer never comes
            return 200, {}, b"{"
        return 200, {}, json.dumps({"message": {"content": '{"a": "yes"}'}}).encode()

    stand_in, url = _stand_in(handle)
    ask = lambda: ai.chat("Copy.", "Say yes.", {"type": "object"})         # noqa: E731
    saved = (ai.FOLDER, ai.SHARED, ai.GIVEN, ai.subprocess, ai._free_port)
    try:
        with tempfile.TemporaryDirectory() as tmp:
            tmp = Path(tmp)
            # Our folder's name is not ASCII, so the note must escape it.
            ai.FOLDER, ai.SHARED, ai.GIVEN = tmp / ai.APP / "ai-कागज़", tmp / "Shared AI", ""
            ai.subprocess = types.SimpleNamespace(**{k: getattr(subprocess, k) for k in (
                "DEVNULL", "TimeoutExpired", "SubprocessError", "run")}, Popen=Runtime)
            ai._free_port = lambda: int(url.rsplit(":", 1)[1])
            bridge = tmp / "Invoice Bridge"
            (bridge / "ollama").mkdir(parents=True)
            (bridge / "ollama" / "ollama.exe").write_bytes(b"the bridge's runtime")
            (bridge / "ollama-models").mkdir()
            (bridge / "ollama-models" / "MODEL").write_text(ai.MODEL + "\r\n", encoding="ascii")
            ai.SHARED.mkdir()
            note = {"runtime": str(bridge / "ollama" / "ollama.exe"), "models": str(bridge / "ollama-models"),
                    "model": ai.MODEL, "ollama": ai.OLLAMA_VERSION}
            for name, listed in (("Invoice Bridge", note), ("Older app", {**note, "ollama": "v0.30.0"}),
                                 ("Relative", {**note, "runtime": "ollama\\ollama.exe"}),
                                 (ai.APP, note)):                    # our own name is never borrowed from
                (ai.SHARED / f"{name}.json").write_text(json.dumps(listed), encoding="ascii")
            (ai.SHARED / "Half written.json").write_text('{"runtime": ', encoding="ascii")
            files = lambda: sorted((str(p), p.stat().st_size, p.stat().st_mtime_ns)   # noqa: E731
                                   for p in bridge.rglob("*"))
            theirs = files()

            assert [c["app"] for c in ai.others()] == ["Invoice Bridge"]
            s = ai.status()
            assert s["status"] == "ready" and s["shared"] == {"app": "Invoice Bridge", "folder": str(bridge)}, s
            assert ask() == {"a": "yes"}
            run = runtimes[-1]
            assert run.args == [str(bridge / "ollama" / "ollama.exe"), "serve"], run.args
            assert run.env["OLLAMA_MODELS"] == str(bridge / "ollama-models") and run.env["OLLAMA_NOPRUNE"] == "1"
            assert run.env["OLLAMA_HOST"].startswith("127.0.0.1:") and run.kw["creationflags"] & 0x40   # idle
            stop_it["now"] = True
            try:
                ask()
                raise AssertionError("a runtime stopped part-way was taken for a bad answer")
            except ai.Stopped:
                pass
            stop_it["now"] = False
            assert ask() == {"a": "yes"} and len(runtimes) == 2, "the runtime was not started again"
            ai.announce()
            assert not (ai.SHARED / f"{ai.APP}.json").exists(), "a note was left for a copy we do not have"
            assert files() == theirs, "the bridge's copy was written to"

            ai.runtime_dir().mkdir(parents=True)                     # our own copy, complete
            (ai.runtime_dir() / "ollama.exe").write_bytes(b"our runtime")
            ai.models_dir().mkdir()
            (ai.models_dir() / "MODEL").write_text(ai.MODEL, encoding="ascii")
            ai.announce()
            raw = (ai.SHARED / f"{ai.APP}.json").read_bytes()
            assert json.loads(raw.decode("ascii")) == {"runtime": str(ai.runtime_dir() / "ollama.exe"),
                                                        "models": str(ai.models_dir()), "model": ai.MODEL,
                                                        "ollama": ai.OLLAMA_VERSION}
            s = ai.status()
            assert s["shared"] is None and s["also"] == {"app": "Invoice Bridge", "folder": str(bridge)}, s
            assert ask() == {"a": "yes"} and runtimes[-1].args[0] == str(ai.runtime_dir() / "ollama.exe")
            assert runtimes[-2].code is not None, "the bridge's runtime was left running beside ours"

            # Programs running from our copy (the bridge may run it) are stopped, and only those.
            ping = Path(os.environ["SystemRoot"]) / "System32" / "PING.EXE"
            running = {}
            for who, exe in (("ours", ai.runtime_dir() / "runner.exe"), ("theirs", bridge / "ollama" / "runner.exe")):
                shutil.copy2(ping, exe)
                running[who] = subprocess.Popen([str(exe), "-n", "30", "127.0.0.1"], creationflags=0x08000000,
                                                stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL)
            swept, sweep = [], ai._stop_all_in
            ai._stop_all_in = lambda folder: (swept.append((folder, (ai.SHARED / f"{ai.APP}.json").exists())),
                                              sweep(folder))
            try:
                ai.remove()
            finally:
                ai._stop_all_in = sweep
                alive = running["theirs"].poll() is None
                running["theirs"].kill()
                running["theirs"].wait()
            assert swept == [(ai.FOLDER, False)], swept              # the note went first
            assert running["ours"].wait(timeout=10) is not None, "a program running from our copy was left"
            assert alive, "a program outside our copy was stopped"
            assert not ai.FOLDER.exists() and not (ai.SHARED / f"{ai.APP}.json").exists()
            s = ai.status()
            assert s["status"] == "ready" and s["shared"]["app"] == "Invoice Bridge", s

            (bridge / "ollama-models" / "MODEL").unlink()              # the bridge is uninstalled
            assert ai.status()["status"] == "absent" and not ai.installed()    # its own download is offered again
    finally:
        ai.stop()
        ai.FOLDER, ai.SHARED, ai.GIVEN, ai.subprocess, ai._free_port = saved
        ai._set(status="absent", done=0, total=0, step="", message="")
        stand_in.shutdown()
