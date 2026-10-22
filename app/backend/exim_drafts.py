"""
Drafted documents for the Exports module: the papers that fix what the
register found, filled in from it.

- The SB005 pack for the port (request letter, Annexure A, challan schedule)
- The Table 9A worksheet for corrections made in GSTR-1
- The SB003 undertaking for a bill filed under another GSTIN of the same PAN
- The declaration to the bank for bills up to the declaration limit
- The request to the bank for more time to bring the money in
- The list of shipping bills to look up on ICEGATE
- The register in the exporter's own sheet layout

Every one is a draft: the reasons, the challan number and the signature are
left for the exporter. Letters are Word files written directly - a .docx is a
zip of three XML parts - so no library ships for them; schedules are Excel.
"""

import io
import re
import zipfile
from collections import defaultdict
from datetime import date, timedelta
from xml.sax.saxutils import escape

import openpyxl
from openpyxl.utils import get_column_letter

from app.backend.common import r2
from app.backend.excel_generator import CURRENCY, DATE_FMT
from app.backend.exports_recon import _sheet
from app.backend.explain import _day, _rs

# Only the codes whose names are certain; any other port is named by its code.
PORT_NAMES = {"INNSA1": "Nhava Sheva (JNPT)", "INMUN1": "Mundra", "INBOM1": "Mumbai", "INMAA1": "Chennai"}


def _fc(v, cur):
    return f"{cur or ''} {abs(v or 0):,.2f}".strip()


def _fy_end(iso):
    """30 November after the financial year the date falls in - the last
    date for Table 9A."""
    y, m = int(iso[:4]), int(iso[5:7])
    return f"30 Nov {y + 1 if m >= 4 else y}"


# ---------------------------------------------------------------------------
# A Word document, written directly
# ---------------------------------------------------------------------------

FONT = '<w:rFonts w:ascii="Arial" w:hAnsi="Arial" w:cs="Arial"/>'


def _run(text, bold=False, size=10.5):
    return (f'<w:r><w:rPr>{FONT}{"<w:b/>" if bold else ""}<w:sz w:val="{int(size * 2)}"/></w:rPr>'
            f'<w:t xml:space="preserve">{escape(str(text))}</w:t></w:r>')


def _p(text="", bold=False, size=10.5, align=None, after=120):
    lines = str(text).split("\n")
    runs = '<w:r><w:br/></w:r>'.join(_run(line, bold, size) for line in lines)
    jc = f'<w:jc w:val="{align}"/>' if align else ""
    return f'<w:p><w:pPr>{jc}<w:spacing w:after="{after}"/></w:pPr>{runs}</w:p>'


def _page_break():
    return '<w:p><w:r><w:br w:type="page"/></w:r></w:p>'


def _table(headers, rows, widths, money=()):
    edge = "".join(f'<w:{e} w:val="single" w:sz="4" w:space="0" w:color="555555"/>'
                   for e in ("top", "left", "bottom", "right", "insideH", "insideV"))

    def cell(text, w, bold=False, right=False):
        jc = '<w:jc w:val="right"/>' if right else ""
        return (f'<w:tc><w:tcPr><w:tcW w:w="{w}" w:type="dxa"/></w:tcPr><w:p><w:pPr>{jc}'
                f'<w:spacing w:after="0"/></w:pPr>{_run(text, bold, 8.5)}</w:p></w:tc>')

    head = "".join(cell(h, w, True) for h, w in zip(headers, widths))
    body = "".join("<w:tr>" + "".join(cell(v, w, right=i in money) for i, (v, w) in enumerate(zip(r, widths)))
                   + "</w:tr>" for r in rows)
    grid = "".join(f'<w:gridCol w:w="{w}"/>' for w in widths)
    return (f'<w:tbl><w:tblPr><w:tblW w:w="{sum(widths)}" w:type="dxa"/><w:tblBorders>{edge}</w:tblBorders>'
            f'<w:tblCellMar><w:left w:w="60" w:type="dxa"/><w:right w:w="60" w:type="dxa"/></w:tblCellMar>'
            f'</w:tblPr><w:tblGrid>{grid}</w:tblGrid><w:tr>{head}</w:tr>{body}</w:tbl>' + _p(after=60))


def docx(body):
    document = ('<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
                '<w:document xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"><w:body>'
                f'{body}<w:sectPr><w:pgSz w:w="11906" w:h="16838"/><w:pgMar w:top="1134" w:right="1021" '
                'w:bottom="1134" w:left="1134" w:header="567" w:footer="567" w:gutter="0"/></w:sectPr>'
                '</w:body></w:document>')
    types = ('<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
             '<Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types">'
             '<Default Extension="rels" ContentType="application/vnd.openxmlformats-package.relationships+xml"/>'
             '<Default Extension="xml" ContentType="application/xml"/>'
             '<Override PartName="/word/document.xml" ContentType="application/'
             'vnd.openxmlformats-officedocument.wordprocessingml.document.main+xml"/></Types>')
    rels = ('<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
            '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
            '<Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/'
            'relationships/officeDocument" Target="word/document.xml"/></Relationships>')
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as z:
        z.writestr("[Content_Types].xml", types)
        z.writestr("_rels/.rels", rels)
        z.writestr("word/document.xml", document)
    return buf.getvalue()


def _letter(entity, today, to, subject, paras, table=None, after_table=(), enclosures=()):
    name = entity.get("name") or "[Your business name]"
    ids = "   ".join(x for x in (f"GSTIN: {entity['gstin']}" if entity.get("gstin") else "",
                                f"IEC: {entity['iec']}" if entity.get("iec") else "") if x)
    out = [_p(name, bold=True, size=13, after=0)]
    if entity.get("address"):
        out.append(_p(entity["address"], size=9.5, after=0))
    if ids:
        out.append(_p(ids, size=9.5, after=0))
    out += [_p(f"Date: {_day(today)}", align="right", after=240), _p("To,\n" + "\n".join(to), after=240),
            _p(f"Subject: {subject}", bold=True), _p("Respected Sir / Madam,")]
    out += [_p(x) for x in paras]
    if table:
        out.append(_table(*table))
    out += [_p(x) for x in after_table]
    if enclosures:
        out.append(_p("Enclosures:\n" + "\n".join(f"{i}. {e}" for i, e in enumerate(enclosures, 1)), size=9.5))
    out += [_p("Thanking you,\nYours faithfully,", after=0), _p(f"For {name}", bold=True, after=720),
            _p(f"{entity.get('signatory') or '[Name]'}\nAuthorised signatory", after=0)]
    return "".join(out)


def _xlsx(fill):
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    fill(wb)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _zip(files):
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as z:
        for name, data in files:
            z.writestr(name, data)
    return buf.getvalue()


def _rows_for(view, status, keep=lambda rec: True):
    by_key = {r["key"]: r for r in view["invoices"]}
    return [(rec, by_key[rec["key"]]) for rec in view["all_records"]
            if rec["status"] == status and rec.get("key") in by_key and keep(rec)]


# ---------------------------------------------------------------------------
# SB005 through the port
# ---------------------------------------------------------------------------

def _sb005_lines(view):
    out = []
    for rec, row in _rows_for(view, "sb005", lambda r: r.get("fix") == "port"):
        sb, gst = row["sb"], row["gst"] or {}
        out.append({"port": sb["port"], "sb_no": sb["sb_no"], "sb_date": sb["sb_date"], "sb_inv": sb["inv_no"],
                    "sb_inv_date": sb["inv_date"], "gst_inv": gst.get("inum"), "gst_inv_date": gst.get("idt"),
                    "igst": sb.get("igst_amount") or gst.get("iamt") or 0})
    return out


def sb005_pack(view, entity, today, fee=1000):
    lines = _sb005_lines(view)
    by_port = defaultdict(list)
    for x in lines:
        by_port[x["port"]].append(x)
    files = []
    for port, xs in sorted(by_port.items()):
        where = f"{PORT_NAMES[port]} ({port})" if port in PORT_NAMES else port
        table = (["S.No.", "SB No.", "SB Date", "Invoice No. in SB", "Invoice Date in SB",
                  "Invoice No. in GSTR-1", "Invoice Date in GSTR-1", "IGST (Rs)"],
                 [[i, x["sb_no"], _day(x["sb_date"]), x["sb_inv"], _day(x["sb_inv_date"]), x["gst_inv"],
                   _day(x["gst_inv_date"]), f"{x['igst']:,.2f}"] for i, x in enumerate(xs, 1)],
                 [500, 1000, 1150, 1500, 1150, 1500, 1150, 1100], (7,))
        body = _letter(
            entity, today,
            ["The Assistant / Deputy Commissioner of Customs", "IGST Refund Section", f"Customs station: {where}"],
            f"Correction of invoice details in {len(xs)} shipping bill(s) through the officer interface - "
            "IGST refund held under error SB005",
            [f"The IGST refund on the shipping bill(s) listed in Annexure A below is held under error SB005 "
             "(invalid invoice number), because the invoice number or date on the shipping bill does not match "
             "the invoice reported in our GSTR-1, Table 6A.",
             "The invoice numbers in GSTR-1 are those of our GST invoices. Annexure A sets out, shipping bill by "
             "shipping bill, the number and date on the shipping bill against the number and date in GSTR-1, for "
             "the same export.",
             "We request that the invoice details be corrected through the officer interface in ICES so that the "
             f"refund can be processed. The fee of Rs {fee:,} per shipping bill, Rs {fee * len(xs):,} for "
             f"{len(xs)} shipping bill(s), is paid by challan no. __________ dated __________ (original enclosed).",
             "We declare that each shipping bill listed and its GST invoice describe the same export, that the "
             "IGST shown was paid, and that no refund has been received for these shipping bills.",
             "Annexure A"],
            table, enclosures=["Annexure A (above, and as a spreadsheet)", "Original challan for the fee",
                               "Copies of the shipping bills and the GST invoices"])
        files.append((f"{port}/Request letter {port}.docx", docx(body)))

        def fill(wb, xs=xs, port=port):
            _sheet(wb, "Annexure A", f"Annexure A - shipping bills at {port} with invoice details to correct",
                   ["S.No.", "SB No.", "SB Date", "Invoice No. in SB", "Invoice Date in SB",
                    "Invoice No. in GSTR-1", "Invoice Date in GSTR-1", "IGST (Rs)"],
                   [[i, x["sb_no"], x["sb_date"], x["sb_inv"], x["sb_inv_date"], x["gst_inv"], x["gst_inv_date"],
                     x["igst"]] for i, x in enumerate(xs, 1)], {3: DATE_FMT, 5: DATE_FMT, 7: DATE_FMT, 8: CURRENCY})
        files.append((f"{port}/Annexure A {port}.xlsx", _xlsx(fill)))

    def challans(wb):
        rows = [[x["port"], x["sb_no"], x["sb_date"], fee] for x in lines]
        _sheet(wb, "Challan schedule", f"Fee of Rs {fee:,} per shipping bill - one challan per customs station",
               ["Customs station", "SB No.", "SB Date", "Fee (Rs)"], rows, {3: DATE_FMT, 4: CURRENCY}, sum_cols=[4])
    files.append(("Challan schedule.xlsx", _xlsx(challans)))
    return _zip(files)


# ---------------------------------------------------------------------------
# Table 9A: corrections made in GSTR-1
# ---------------------------------------------------------------------------

def _9a_lines(view):
    out = []
    for status in ("sb001", "sb005", "type_differs", "igst_differs", "gst_twice"):
        for rec, row in _rows_for(view, status, lambda r: r["status"] != "sb005" or r.get("fix") in ("9a", "check")):
            sb, g = row["sb"] or {}, row["gst"] or {}
            line = {"period": g.get("doc_name"), "orig_no": g.get("inum"), "orig_date": g.get("idt"),
                    "new_no": g.get("inum"), "new_date": g.get("idt"), "type": g.get("exp_typ"),
                    "port": g.get("sbpcode"), "sb_no": g.get("sbnum"), "sb_date": g.get("sbdt"),
                    "txval": g.get("txval"), "igst": g.get("iamt"), "amend_by": _fy_end(g.get("idt") or sb.get("sb_date"))}
            if status == "sb001":
                line.update(port=sb.get("port"), sb_no=sb.get("sb_no"), sb_date=sb.get("sb_date"),
                            change="Shipping bill number, date and port, as on the shipping bill")
            elif status == "sb005":
                line.update(new_no=sb.get("inv_no"), new_date=sb.get("inv_date"),
                            change="Invoice number and date, as on your invoice and the shipping bill"
                            if rec.get("fix") == "9a" else
                            "Invoice number or date - only if Table 6A is the wrong side (check your invoice first)")
            elif status == "type_differs":
                wpay = sb.get("scheme") == "WPAY"
                line.update(type="WPAY" if wpay else "WOPAY", igst=sb.get("igst_amount") if wpay else 0,
                            change="Export type (with or without IGST), as on the shipping bill")
            elif status == "igst_differs":
                line.update(igst=sb.get("igst_amount"), change="IGST, as on the shipping bill - if the shipping bill is right")
            else:
                line.update(change="Remove the duplicate line")
            out.append(line)
    return out


def table9a_workbook(view):
    lines = _9a_lines(view)
    add = [row for _, row in _rows_for(view, "not_in_gst")]

    def fill(wb):
        _sheet(wb, "Table 9A", "Table 6A lines to amend through Table 9A - one line per invoice",
               ["S.No.", "Original return", "Original invoice no.", "Original invoice date", "Revised invoice no.",
                "Revised invoice date", "Export type", "Port code", "Shipping bill no.", "Shipping bill date",
                "Taxable value (Rs)", "IGST (Rs)", "What changes", "Amend by"],
               [[i, x["period"], x["orig_no"], x["orig_date"], x["new_no"], x["new_date"], x["type"], x["port"],
                 x["sb_no"], x["sb_date"], x["txval"], x["igst"], x["change"], x["amend_by"]]
                for i, x in enumerate(lines, 1)],
               {4: DATE_FMT, 6: DATE_FMT, 10: DATE_FMT, 11: CURRENCY, 12: CURRENCY})
        _sheet(wb, "Add to Table 6A", "Exports missing from GSTR-1 - report them in the current month's Table 6A",
               ["S.No.", "Invoice no.", "Invoice date", "Export type", "Port code", "Shipping bill no.",
                "Shipping bill date", "Invoice value", "Currency", "IGST (Rs)"],
               [[i, r["inv_no"], r["inv_date"], r["sb"]["scheme"] if r["sb"]["scheme"] == "WPAY" else "WOPAY",
                 r["sb"]["port"], r["sb"]["sb_no"], r["sb"]["sb_date"], r["value_fc"], r["currency"],
                 r["sb"].get("igst_amount")] for i, r in enumerate(add, 1)],
               {3: DATE_FMT, 7: DATE_FMT, 8: CURRENCY, 10: CURRENCY})
    return _xlsx(fill)


# ---------------------------------------------------------------------------
# SB003: another registration of the same PAN
# ---------------------------------------------------------------------------

def _sb003_lines(view):
    gstin = (view.get("entity") or {}).get("gstin")
    out = []
    for rec, row in _rows_for(view, "sb003"):
        sb = row["sb"]
        if gstin and sb.get("gstin") and sb["gstin"][2:12] == gstin[2:12]:     # the same PAN only
            out.append({"port": sb["port"], "sb_no": sb["sb_no"], "sb_date": sb["sb_date"], "inv": row["inv_no"],
                        "sb_gstin": sb["gstin"], "gstin": gstin, "igst": sb.get("igst_amount") or 0})
    return out


def sb003_undertaking(view, entity, today):
    lines = _sb003_lines(view)
    by_port = defaultdict(list)
    for x in lines:
        by_port[x["port"]].append(x)
    parts = []
    for port, xs in sorted(by_port.items()):
        where = f"{PORT_NAMES[port]} ({port})" if port in PORT_NAMES else port
        sb_gstins = sorted({x["sb_gstin"] for x in xs})
        parts.append(_letter(
            entity, today,
            ["The Assistant / Deputy Commissioner of Customs", "IGST Refund Section", f"Customs station: {where}"],
            f"Undertaking for {len(xs)} shipping bill(s) filed under another GSTIN of the same PAN - "
            "IGST refund held under error SB003",
            [f"The shipping bill(s) listed below were filed under our GSTIN {', '.join(sb_gstins)}, while the "
             f"export invoices were reported in the GSTR-1 of our GSTIN {xs[0]['gstin']}. Both registrations "
             f"belong to the same PAN, {xs[0]['gstin'][2:12]}, so ICEGATE returns error SB003 (GSTIN mismatch) "
             "and the refund is held.",
             f"We request that the refund be processed through the officer interface against the return filed "
             f"under GSTIN {xs[0]['gstin']}."],
            (["S.No.", "SB No.", "SB Date", "Invoice", "GSTIN on SB", "GSTIN in GSTR-1", "IGST (Rs)"],
             [[i, x["sb_no"], _day(x["sb_date"]), x["inv"], x["sb_gstin"], x["gstin"], f"{x['igst']:,.2f}"]
              for i, x in enumerate(xs, 1)], [500, 1050, 1150, 1600, 1850, 1850, 1250], (6,)),
            [f"We undertake that no refund has been claimed, or will be claimed, for these exports under GSTIN "
             f"{', '.join(sb_gstins)}, and that we will repay with interest any refund later found not to be due."],
            enclosures=["Copies of the shipping bills", "GSTR-1 (Table 6A) of the return showing these invoices",
                        "Registration certificates of both GSTINs"]))
    return docx(_page_break().join(parts))


# ---------------------------------------------------------------------------
# Letters to the bank: the declaration and the extension
# ---------------------------------------------------------------------------

def _bank_to(view):
    ads = [r["sb"].get("ad_code") for r in view["invoices"] if r["sb"] and r["sb"].get("ad_code")]
    ad = max(set(ads), key=ads.count) if ads else None
    return ["The Branch Manager", "[Your AD bank and branch]", f"AD code: {ad}" if ad else "AD code: [    ]"]


def _declaration_lines(view, limit):
    out = []
    for r in view["invoices"]:
        due = r.get("due") or {}
        if not due or not r["sb"]:
            continue
        if due.get("short_by_charges"):
            short = due["short_by_charges"]
            out.append({"row": r, "kind": "Reduction", "amount_fc": short, "inr": r2(short * (r["rate"] or 0)),
                        "reason": "Bank charges deducted abroad"})
        elif due["state"].startswith("overdue") and 0 < due["outstanding_inr"] <= limit:
            out.append({"row": r, "kind": "Not realised (if it will not come)", "amount_fc": due["outstanding_fc"],
                        "inr": due["outstanding_inr"], "reason": ""})
    return out


def declaration(view, entity, today, limit=1000000, start="2026-10-01"):
    lines = _declaration_lines(view, limit)
    table = (["S.No.", "SB No. / Date", "Port", "Invoice", "Currency", "Invoice value", "To declare",
              "Rs equivalent", "Type", "Reason"],
             [[i, f"{x['row']['sb']['sb_no']} / {_day(x['row']['sb']['sb_date'])}", x["row"]["sb"]["port"],
               x["row"]["inv_no"], x["row"]["currency"], f"{x['row']['value_fc'] or 0:,.2f}",
               f"{x['amount_fc']:,.2f}", f"{x['inr']:,.0f}", x["kind"], x["reason"] or "[reason]"]
              for i, x in enumerate(lines, 1)],
             [420, 1300, 750, 1150, 650, 1000, 950, 950, 1150, 1400], (5, 6, 7))
    body = _letter(
        entity, today, _bank_to(view),
        f"Declaration for closing or reducing {len(lines)} export bill(s) of up to Rs {limit:,} in EDPMS",
        [f"Under the export regulations in force from {_day(start)}, an export bill of up to Rs {limit:,} can be "
         "closed or reduced in EDPMS on the exporter's declaration.",
         "We declare that the amounts shown against the bills below will not be realised, for the reasons given "
         "against each, and request that the bills be closed or reduced in EDPMS accordingly."],
        table,
        ["We confirm that the amounts not realised are not due to any action of ours to keep the proceeds abroad, "
         "and that we will bring in any amount later received."],
        enclosures=["Copies of the shipping bills", "Correspondence with the buyers, where there is any"])

    def fill(wb):
        _sheet(wb, "Declaration", f"Bills of up to Rs {limit:,} to close or reduce in EDPMS",
               ["SB No.", "SB Date", "Port", "Invoice", "Buyer", "Currency", "Invoice value", "To declare",
                "Rs equivalent", "Type", "Reason"],
               [[x["row"]["sb"]["sb_no"], x["row"]["sb"]["sb_date"], x["row"]["sb"]["port"], x["row"]["inv_no"],
                 x["row"]["buyer"], x["row"]["currency"], x["row"]["value_fc"], x["amount_fc"], x["inr"], x["kind"],
                 x["reason"]] for x in lines], {2: DATE_FMT, 7: CURRENCY, 8: CURRENCY, 9: CURRENCY})
    return _zip([("Declaration to the bank.docx", docx(body)), ("Declaration schedule.xlsx", _xlsx(fill))])


def _extension_lines(view):
    return [r for r in view["invoices"] if r.get("due") and r["due"]["state"] in ("overdue", "overdue_year", "due_soon")]


def extension(view, entity, today):
    rows = _extension_lines(view)
    want = lambda r: date.fromisoformat(r["due"]["date"]) + timedelta(days=180)
    table = (["S.No.", "SB No. / Date", "Invoice", "Buyer", "Outstanding", "Due date", "Extension to", "Reason"],
             [[i, f"{r['sb']['sb_no']} / {_day(r['sb']['sb_date'])}", r["inv_no"], r["buyer"] or "",
               _fc(r["due"]["outstanding_fc"], r["currency"]), _day(r["due"]["date"]), _day(want(r).isoformat()),
               "[reason]"] for i, r in enumerate(rows, 1)],
             [420, 1500, 1150, 1700, 1300, 1050, 1050, 1550], (4,))
    body = _letter(
        entity, today, _bank_to(view),
        f"Request for more time to realise the export proceeds of {len(rows)} bill(s)",
        ["We request an extension of the time for realising the export proceeds of the bills listed below, to the "
         "dates shown, for the reasons given against each.",
         "The buyers have confirmed the amounts are payable, and we undertake to bring the proceeds in within the "
         "extended time."],
        table, enclosures=["Copies of the shipping bills", "The buyers' confirmation of payment"])

    def fill(wb):
        _sheet(wb, "Extension", "Bills to ask the bank for more time on",
               ["SB No.", "SB Date", "Port", "Invoice", "Buyer", "Currency", "Outstanding", "Due date",
                "Extension to", "Reason"],
               [[r["sb"]["sb_no"], r["sb"]["sb_date"], r["sb"]["port"], r["inv_no"], r["buyer"], r["currency"],
                 r["due"]["outstanding_fc"], r["due"]["date"], want(r).isoformat(), ""] for r in rows],
               {2: DATE_FMT, 7: CURRENCY, 8: DATE_FMT, 9: DATE_FMT})
    return _zip([("Extension request to the bank.docx", docx(body)), ("Extension schedule.xlsx", _xlsx(fill))])


# ---------------------------------------------------------------------------
# ICEGATE look-ups, and the register in the exporter's own layout
# ---------------------------------------------------------------------------

STATE_WORDS = {"held": "held", "pending": "not received", "unknown": "no status on file", "waiting": "not due yet"}


def _lookup_lines(view):
    return [r for r in view["invoices"] if r.get("icegate")]


def lookup_workbook(view):
    rows = _lookup_lines(view)
    pages = (view.get("rules") or {}).get("icegate") or {}

    def fill(wb):
        ws = _sheet(wb, "Look up on ICEGATE", "Shipping bills whose refund or drawback is not confirmed as received",
                    ["SB No.", "SB Date", "Port", "Invoice", "IGST (Rs)", "IGST", "Code known", "Drawback (Rs)",
                     "Drawback", "What to look up", "On ICEGATE"],
                    [[r["sb"]["sb_no"], r["sb"]["sb_date"], r["sb"].get("port"), r["inv_no"],
                      r["refund"]["igst"]["due"] or None,
                      STATE_WORDS.get(r["refund"]["igst"]["state"], r["refund"]["igst"]["state"]),
                      r["refund"]["igst"].get("code") or "", r["refund"]["dbk"]["due"] or None,
                      STATE_WORDS.get(r["refund"]["dbk"]["state"], r["refund"]["dbk"]["state"]),
                      " and ".join(x for x, s in (("IGST refund status", r["refund"]["igst"]["state"]),
                                                  ("drawback status", r["refund"]["dbk"]["state"]))
                                   if s in ("held", "pending", "unknown")),
                      "Shipping bill enquiry" if pages.get("sb_enquiry") else ""]
                     for r in rows], {2: "dd-mm-yyyy", 5: CURRENCY, 8: CURRENCY})
        # The date shows as ICEGATE's form asks for it; each link opens the enquiry.
        ws.cell(2, 1, "Each link opens ICEGATE's public shipping bill enquiry (no login): choose the port, type the SB "
                      "number and date as shown here, and the captcha. The IGST refund status needs your ICEGATE login: "
                      f"{pages.get('igst_path') or 'Services > Enquiries'}.")
        for n in range(len(rows)):
            cell = ws.cell(4 + n, 11)
            if pages.get("sb_enquiry"):
                cell.hyperlink, cell.style = pages["sb_enquiry"], "Hyperlink"
    return _xlsx(fill)


OWN_HEADERS = ["INV.NO.", "INV.DATE", "INV. AMOUNT", "CONTAINER NO.", "ICD", "S.B.NO.", "S.B.DATE", "GST AMOUNT",
               "GST AMO.RECD", "DIFF", "DATE OF RECD", "D.BACK AMOUNT", "DATE OF RECD", "ERROR CODE", "PAYMENT",
               "BANK", "DEL. TO BUYER"]


def own_format_workbook(view):
    """The register in the layout the exporter already keeps: one tab per
    financial year, their columns, filled in from every document on file."""
    years = defaultdict(list)
    for r in view["invoices"]:
        sb, own = r["sb"] or {}, (r["own"] or [{}])[0]
        when = sb.get("sb_date") or r["inv_date"]
        if not when:
            continue
        y, m = int(when[:4]), int(when[5:7])
        fy = f"{y if m >= 4 else y - 1}-{str((y if m >= 4 else y - 1) + 1)[2:]}"
        ig, db = (r.get("refund") or {}).get("igst") or {}, (r.get("refund") or {}).get("dbk") or {}
        due = r.get("due") or {}
        years[fy].append([
            r["inv_no"], r["inv_date"], r["value_fc"] if r["value_fc"] else own.get("inv_amount"),
            ", ".join(sb.get("containers") or []) or own.get("container"), own.get("icd") or sb.get("port"),
            sb.get("sb_no"), sb.get("sb_date"), ig.get("due") or None, ig.get("recd"),
            r2(ig["due"] - ig["recd"]) if ig.get("recd") is not None and ig.get("due") else None, ig.get("recd_date"),
            db.get("due") or None, db.get("recd_date"),
            ig.get("code") if ig.get("state") == "held" else "",
            "RECD" if due.get("state") == "paid" or any(o.get("payment_recd") for o in r["own"]) else "",
            own.get("bank"), own.get("docs")])

    def fill(wb):
        for fy in sorted(years):
            rows = sorted(years[fy], key=lambda x: (x[6] or x[1] or "", x[0] or ""))
            ws = wb.create_sheet(fy)
            ws.append(OWN_HEADERS)
            for row in rows:
                ws.append([date.fromisoformat(v) if isinstance(v, str) and re.fullmatch(r"\d{4}-\d{2}-\d{2}", v)
                           else v for v in row])
            for col in "BGKM":
                for cell in ws[col][1:]:
                    cell.number_format = "dd.mm.yyyy"
            for col in "CHIJL":
                for cell in ws[col][1:]:
                    cell.number_format = "#,##0.00"
            ws.freeze_panes = "A2"
            for i, w in enumerate([12, 11, 12, 16, 9, 10, 11, 12, 12, 9, 11, 12, 11, 10, 9, 11, 13], start=1):
                ws.column_dimensions[get_column_letter(i)].width = w
        if not years:
            wb.create_sheet("Register").append(OWN_HEADERS)
    return _xlsx(fill)


# ---------------------------------------------------------------------------
# What is ready to draft
# ---------------------------------------------------------------------------

def summary(view, limit=1000000):
    sb005 = _sb005_lines(view)
    return {
        "sb005": {"count": len(sb005), "ports": sorted({x["port"] for x in sb005}), "fee": 1000 * len(sb005),
                  "check": len(_rows_for(view, "sb005", lambda r: r.get("fix") == "check"))},
        "table9a": {"count": len(_9a_lines(view)), "add": len(_rows_for(view, "not_in_gst"))},
        "sb003": {"count": len(_sb003_lines(view))},
        "declaration": {"count": len(_declaration_lines(view, limit))},
        "extension": {"count": len(_extension_lines(view))},
        "lookup": {"count": len(_lookup_lines(view))},
    }
