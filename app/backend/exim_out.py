"""
Files out of the Exports module: the register as a workbook, and DGFT's bulk
eBRC sheet filled in from it.

The bulk sheet's first tab carries only DGFT's columns, headers in the first
row, so it can go into the portal's template as it is; the checks sit on a tab
of their own for the same reason.
"""

import io
from datetime import datetime

import openpyxl

from app.backend.excel_generator import (
    ALIGN_LEFT, CURRENCY, DATE_FMT, FILL_CLEAR, FILL_HEADER, FILL_RISK, FILL_WARN,
    FONT_DATA, FONT_TBL_HDR, auto_fit_columns,
)
from app.backend.exports_recon import _sheet

CHIP_FILL = {"ok": FILL_CLEAR, "marked": FILL_CLEAR, "missing": FILL_RISK, "problem": FILL_RISK,
             "look": FILL_WARN, "pending": FILL_WARN}
CHIPS = (("sb", "Shipping bill"), ("bl", "Bill of lading"), ("gst", "GSTR-1 (Table 6A)"),
         ("paid", "Payment"), ("ebrc", "eBRC"))
DID = {"added": "Added", "posted": "Checked and posted", "corrected": "Corrected", "removed": "Removed",
       "payment": "Payment placed", "marked": "Your answer", "decision": "Decision"}


def _dgft_day(iso):
    """DGFT's upload wants 05-Jan-2026, as text."""
    try:
        return datetime.strptime(str(iso)[:10], "%Y-%m-%d").strftime("%d-%b-%Y")
    except ValueError:
        return iso or ""


def register_workbook(payload):
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    name = (payload.get("entity") or {}).get("name") or "Exports"

    rows = []
    for r in payload["invoices"]:
        due, sb, bl = r.get("due") or {}, r.get("sb") or {}, (r.get("bls") or [{}])[0]
        rows.append([r["inv_no"], r.get("inv_date"), r.get("buyer"), r.get("currency"), r.get("value_fc"),
                     sb.get("sb_no"), sb.get("sb_date"), sb.get("port"), bl.get("bl_no"), bl.get("sob_date"),
                     due.get("received_fc"), due.get("outstanding_fc"), due.get("date")]
                    + [r["chips"][k]["text"] for k, _ in CHIPS])
    headers = ["Invoice", "Invoice date", "Buyer", "Currency", "Invoice value", "Shipping bill", "SB date",
               "Port", "Bill of lading", "Shipped on board", "Received", "Still to come", "Payment due by",
               *[label for _, label in CHIPS]]
    ws = _sheet(wb, "Register", f"{name} - export register as at {payload['as_at']}", headers, rows,
                {2: DATE_FMT, 5: CURRENCY, 7: DATE_FMT, 10: DATE_FMT, 11: CURRENCY, 12: CURRENCY, 13: DATE_FMT})
    first = len(headers) - len(CHIPS) + 1
    for i, r in enumerate(payload["invoices"]):
        for j, (k, _) in enumerate(CHIPS):
            fill = CHIP_FILL.get(r["chips"][k]["state"])
            if fill:
                ws.cell(4 + i, first + j).fill = fill

    _sheet(wb, "Needs a look", "Every open entry, with why it is open and what to do",
           ["", "Invoice", "What", "What was found", "What to do", "Amount", "Date"],
           [["Serious" if e["tone"] == "serious" else "Look", e.get("inv_no") or "", e["why"]["title"],
             e["why"]["what"], e["why"].get("todo") or "", e.get("amount_text") or "", e.get("date")]
            for e in payload["all_records"]],
           {7: DATE_FMT})

    _sheet(wb, "Money in", "Every inward remittance and what it is set against",
           ["Received", "Remitter", "IRM", "Bank reference", "Currency", "Amount", "Not yet set against an invoice",
            "Purpose code", "Set against", "From"],
           [[m["date"], m["remitter"], m["irm_no"], m["ref"], m["currency"], m["amount_fc"], m["available_fc"],
             m["purpose_code"], "; ".join(f"{x['inv_no']} {x['amount_fc']:,.2f} ({x['how']})" for x in m["mapped"]),
             " and ".join(m["sources"])] for m in payload["remittances"] if not m["cancelled"]],
           {1: DATE_FMT, 6: CURRENCY, 7: CURRENCY})

    _sheet(wb, "eBRCs", "eBRCs on file",
           ["eBRC", "Date", "IRM", "Invoice", "Shipping bill", "Currency", "Realised", "Status"],
           [[e.get("ebrc_no"), e.get("ebrc_date"), e.get("irm_no"), e.get("gst_inv_no") or e.get("bill_no"),
             e.get("sb_no"), e.get("currency"), e.get("realised_fc"), e.get("status")]
            for e in payload["ebrcs"]],
           {2: DATE_FMT, 7: CURRENCY})
    if payload.get("audit"):
        _sheet(wb, "History", "Everything done to this register, newest first",
               ["When", "By", "What", "Details", "Changes (before → after)"],
               [[a["at"].replace("T", " "), a.get("by") or "", DID.get(a["what"], a["what"]), a["detail"],
                 "; ".join(f"{c['field']}: {c['old'] or 'blank'} → {c['new'] or 'blank'}" for c in a["changes"])]
                for a in payload["audit"]], {})
    buf = io.BytesIO()
    wb.save(buf)
    return buf


DGFT_COLUMNS = ["S.No.", "Branch Code", "IFSC Code", "AD Code", "IRM Number", "IRM Date", "IRM Currency",
                "Purpose Code", "IRM Amount Available", "Shipping Bill No.", "Shipping Bill Date", "Port Code",
                "Bill No.", "Amount Mapped", "Vostro", "Deduction", "Avail GST Benefit (Y/N)",
                "GST Invoice No.", "GST Invoice Date"]


def ebrc_bulk_workbook(payload):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "eBRC bulk"
    for c, h in enumerate(DGFT_COLUMNS, start=1):
        cell = ws.cell(1, c, h)
        cell.font, cell.fill, cell.alignment = FONT_TBL_HDR, FILL_HEADER, ALIGN_LEFT
    for i, b in enumerate(payload["bulk_rows"], start=1):
        values = [i, "", b["ifsc"], b["ad_code"], b["irm_no"], _dgft_day(b["irm_date"]), b["currency"],
                  b["purpose_code"], b["available_fc"], b["sb_no"], _dgft_day(b["sb_date"]), b["port"],
                  b["bill_no"], b["mapped_fc"], "N", 0, b["gst_availed"], b["gst_inv_no"],
                  _dgft_day(b["gst_inv_date"])]
        for c, v in enumerate(values, start=1):
            cell = ws.cell(1 + i, c, v if v is not None else "")
            cell.font = FONT_DATA
            if c in (9, 14, 16):
                cell.number_format = "0.00"
    ws.freeze_panes = "A2"
    auto_fit_columns(ws, skip_title_rows=False)

    _sheet(wb, "Row checks", "Each row checked the way DGFT checks it - fix the red ones before uploading",
           ["S.No.", "Invoice", "Remitter", "Ready to upload", "What to fix"],
           [[i, b["bill_no"], b["remitter"], "No" if b["problems"] else "Yes", "; ".join(b["problems"])]
            for i, b in enumerate(payload["bulk_rows"], start=1)], {})
    checks = wb["Row checks"]
    for i, b in enumerate(payload["bulk_rows"]):
        checks.cell(4 + i, 4).fill = FILL_RISK if b["problems"] else FILL_CLEAR
    buf = io.BytesIO()
    wb.save(buf)
    return buf
